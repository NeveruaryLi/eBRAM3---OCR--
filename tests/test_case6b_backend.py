import asyncio
import io
import base64
import json
import unittest
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

import fitz
from docx import Document
from fastapi.testclient import TestClient

from api.pdf_chat import session_metadata, session_results_store, session_store
from app import app
from cases.case6b_service_agreement import (
    Case6BDraftingHandler,
    Case6BSession,
    MaterialRecord,
    _merge_manifest,
    _send_message,
    _validate_agent_fields,
    _apply_profile_evidence_enrichment,
    apply_review_changes,
    begin_analysis_run,
    build_template_preflight,
    build_agent_l_fill_payload,
    build_agent_l_template_payload,
    extract_csv_markdown,
    extract_docx_markdown,
    extract_template_manifest,
    detect_fact_conflicts,
    extract_xlsx_markdown,
    image_to_pdf,
    normalize_agent_conflicts,
    ocr_submission_filename,
    parse_agent_json,
    render_draft_docx,
    request_analysis_cancel,
    validate_material,
    validate_template,
)


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "input_example" / "Usecase 6B" / "Input documents"
TEMPLATE = SAMPLES / "UC6B_0_Format - Service Agreement.docx"
FIXTURES = ROOT / "GPTbots_.bot" / "generated" / "case6b" / "fixtures"


class Case6BCoreTests(unittest.TestCase):
    def test_sample_template_has_37_stable_fields(self):
        template = validate_template(TEMPLATE.read_bytes(), TEMPLATE.name)
        manifest = extract_template_manifest(template.content)
        self.assertEqual(37, len(manifest["fields"]))
        self.assertEqual("p002_f01", manifest["fields"][0]["field_id"])
        self.assertEqual(
            "t000_r000_c001_p003_f01",
            manifest["fields"][-1]["field_id"],
        )
        self.assertNotIn("profile_id", manifest)
        self.assertNotIn("template_defaults", manifest)
        self.assertNotIn("allowed_rewrites", manifest)

    def test_bracket_placeholders_are_detected_for_unknown_templates(self):
        output = io.BytesIO()
        document = Document()
        document.add_heading("Consulting Agreement")
        document.add_paragraph("Between [Provider_Name] and [Client_Name].")
        document.save(output)
        template = validate_template(output.getvalue(), "brackets.docx")
        manifest = extract_template_manifest(template.content)
        self.assertEqual(2, len(manifest["fields"]))
        self.assertEqual("bracket", manifest["fields"][0]["placeholder_type"])
        preflight = build_template_preflight(manifest)
        self.assertEqual(2, preflight["field_count"])
        self.assertTrue(preflight["analysis_ready"])

    def test_chinese_table_service_rows_and_signature_are_classified(self):
        output = io.BytesIO()
        document = Document()
        document.add_paragraph("服務協議：本文件用於記錄雙方服務範圍、價格及正式簽署資料。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "服務項目"
        table.cell(0, 1).text = "服務價格及計費方式"
        table.cell(1, 0).text = "________________"
        table.cell(1, 1).text = "________________"
        document.add_paragraph("簽署人姓名：________________")
        document.save(output)

        template = validate_template(output.getvalue(), "chinese.docx")
        manifest = extract_template_manifest(template.content)

        self.assertEqual("zh_tw", template.language)
        self.assertEqual("zh_tw", manifest["template_language"])
        self.assertEqual(
            ["signature_name", "repeatable_service", "repeatable_price"],
            [field["field_kind"] for field in manifest["fields"]],
        )
        self.assertEqual(
            [None, 1, 1],
            [field["group_index"] for field in manifest["fields"]],
        )

    def test_agent_can_promote_unknown_scalar_to_signature_field(self):
        output = io.BytesIO()
        document = Document()
        document.add_paragraph("Contact [Contact_Name]")
        document.save(output)
        manifest = extract_template_manifest(output.getvalue())
        parsed = {
            "template_language": "en",
            "fields": [
                {
                    "field_id": manifest["fields"][0]["field_id"],
                    "semantic_key": "signatory_name",
                    "label": "Signatory name",
                    "field_kind": "signature_name",
                    "group_key": None,
                    "group_index": None,
                    "required": False,
                    "resolution_policy": "leave_blank",
                }
            ],
        }

        merged = _merge_manifest(manifest, parsed)

        self.assertEqual("signature_name", merged["fields"][0]["field_kind"])
        self.assertEqual(
            manifest["fields"][0]["locator"],
            merged["fields"][0]["locator"],
        )

    def test_party_company_name_is_not_locked_as_personal_signature(self):
        output = io.BytesIO()
        document = Document()
        document.add_paragraph("Client Name: [Client_Name]")
        document.save(output)
        manifest = extract_template_manifest(output.getvalue())
        self.assertEqual("scalar", manifest["fields"][0]["field_kind"])
        parsed = {
            "template_language": "en",
            "fields": [
                {
                    "field_id": manifest["fields"][0]["field_id"],
                    "semantic_key": "client_identity",
                    "label": "Client company name",
                    "field_kind": "scalar",
                    "group_key": None,
                    "group_index": None,
                    "required": True,
                    "resolution_policy": "fill",
                }
            ],
        }
        merged = _merge_manifest(manifest, parsed)
        filled = _validate_agent_fields(
            merged,
            {
                "fields": [
                    {
                        "field_id": manifest["fields"][0]["field_id"],
                        "status": "FILLED",
                        "value": "Acme Limited",
                        "evidence": [
                            {
                                "source": "company.pdf",
                                "fact": "Client: Acme Limited",
                            }
                        ],
                    }
                ]
            },
            filled=True,
        )
        self.assertEqual("FILLED", filled[0]["status"])
        self.assertEqual("Acme Limited", filled[0]["value"])

    def test_filled_repeatable_fields_require_service_action(self):
        manifest = {
            "fields": [
                {
                    "field_id": "p001_f01",
                    "field_kind": "repeatable_service",
                }
            ]
        }
        parsed = {
            "fields": [
                {
                    "field_id": "p001_f01",
                    "status": "FILLED",
                    "value": "Support",
                    "evidence": [{"source": "pricing.xlsx", "fact": "Support"}],
                }
            ]
        }

        with self.assertRaisesRegex(ValueError, "service_action"):
            _validate_agent_fields(manifest, parsed, filled=True)

    def test_repeatable_service_action_matrix_rejects_invalid_values(self):
        manifest = {
            "fields": [
                {
                    "field_id": "p001_f01",
                    "field_kind": "repeatable_service",
                }
            ]
        }
        for status, action in (
            ("NEEDS_CONFIRMATION", "garbage"),
            ("REMOVE", "optional"),
            ("KEEP_BLANK", "included"),
            ("LEAVE_BLANK", "blank"),
        ):
            with self.subTest(status=status, action=action):
                with self.assertRaisesRegex(ValueError, "service_action"):
                    _validate_agent_fields(
                        manifest,
                        {
                            "fields": [
                                {
                                    "field_id": "p001_f01",
                                    "status": status,
                                    "value": "",
                                    "evidence": [],
                                    "service_action": action,
                                }
                            ]
                        },
                        filled=True,
                    )

    def test_repeatable_name_and_price_require_the_same_service_action(self):
        manifest = {
            "fields": [
                {
                    "field_id": "p001_f01",
                    "field_kind": "repeatable_service",
                    "group_index": 1,
                },
                {
                    "field_id": "p001_f02",
                    "field_kind": "repeatable_price",
                    "group_index": 1,
                },
            ]
        }
        parsed = {
            "fields": [
                {
                    "field_id": "p001_f01",
                    "status": "FILLED",
                    "value": "Support",
                    "evidence": [{"source": "scope.xlsx", "fact": "Support"}],
                    "service_action": "included",
                },
                {
                    "field_id": "p001_f02",
                    "status": "FILLED",
                    "value": "HKD 1,000 per month",
                    "evidence": [
                        {"source": "scope.xlsx", "fact": "HKD 1,000 per month"}
                    ],
                    "service_action": "optional",
                },
            ]
        }
        with self.assertRaisesRegex(ValueError, "名称与价格状态不一致"):
            _validate_agent_fields(manifest, parsed, filled=True)

    def test_rejects_service_group_with_more_than_name_price_pair(self):
        local_manifest = {
            "template_language": "en",
            "fields": [
                {
                    "field_id": "p001_f01",
                    "field_kind": "scalar",
                    "locator": "paragraph:1/blank:1",
                },
                {
                    "field_id": "p001_f02",
                    "field_kind": "scalar",
                    "locator": "paragraph:1/blank:2",
                },
                {
                    "field_id": "p001_f03",
                    "field_kind": "scalar",
                    "locator": "paragraph:1/blank:3",
                },
            ],
        }
        parsed = {
            "template_language": "en",
            "fields": [
                {
                    "field_id": "p001_f01",
                    "field_kind": "repeatable_service",
                    "group_key": "services",
                    "group_index": 1,
                },
                {
                    "field_id": "p001_f02",
                    "field_kind": "repeatable_service",
                    "group_key": "services",
                    "group_index": 1,
                },
                {
                    "field_id": "p001_f03",
                    "field_kind": "repeatable_price",
                    "group_key": "services",
                    "group_index": 1,
                },
            ],
        }
        with self.assertRaisesRegex(ValueError, "名称与价格字段"):
            _merge_manifest(local_manifest, parsed)

    def test_rejects_docx_without_placeholders(self):
        output = io.BytesIO()
        document = Document()
        document.add_paragraph("No blank fields")
        document.save(output)
        with self.assertRaisesRegex(ValueError, "下划线"):
            validate_template(output.getvalue(), "plain.docx")

    def test_rejects_unsupported_complex_word_objects(self):
        base = io.BytesIO()
        document = Document()
        document.add_paragraph("Field: ____________________")
        document.save(base)
        markers = {
            "Word 内容控件": b"<w:sdt>",
            "文本框": b"<w:txbxContent>",
            "邮件合并域": b"MERGEFIELD",
            "altChunk": b'<w:altChunk r:id="rId999"/>',
        }
        for expected, marker in markers.items():
            source = zipfile.ZipFile(io.BytesIO(base.getvalue()))
            output = io.BytesIO()
            with source, zipfile.ZipFile(output, "w") as target:
                for item in source.infolist():
                    content = source.read(item.filename)
                    if item.filename == "word/document.xml":
                        content = content.replace(b"</w:body>", marker + b"</w:body>")
                    target.writestr(item, content)
            with self.subTest(expected=expected):
                with self.assertRaisesRegex(ValueError, expected):
                    validate_template(output.getvalue(), "complex.docx")

    def test_rejects_external_docx_preview_resources(self):
        base = io.BytesIO()
        document = Document()
        document.add_paragraph("Field: ____________________")
        document.save(base)
        source = zipfile.ZipFile(io.BytesIO(base.getvalue()))
        output = io.BytesIO()
        relationship = (
            b'<Relationship Id="rId999" '
            b'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            b'relationships/image" Target="https://attacker.invalid/pixel.png" '
            b'TargetMode="External"/>'
        )
        with source, zipfile.ZipFile(output, "w") as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == "word/_rels/document.xml.rels":
                    content = content.replace(
                        b"</Relationships>",
                        relationship + b"</Relationships>",
                    )
                target.writestr(item, content)
        with self.assertRaisesRegex(ValueError, "外部资源"):
            validate_template(output.getvalue(), "external-resource.docx")

    def test_rejects_docx_preview_style_injection(self):
        base = io.BytesIO()
        document = Document()
        document.add_paragraph("Field: ____________________")
        document.save(base)
        source = zipfile.ZipFile(io.BytesIO(base.getvalue()))
        output = io.BytesIO()
        with source, zipfile.ZipFile(output, "w") as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename == "word/styles.xml":
                    replacement = (
                        b'w:styleId="Normal}{body{background:url('
                        b'https://attacker.invalid/pixel)"'
                    )
                    self.assertIn(b'w:styleId="Normal"', content)
                    content = content.replace(
                        b'w:styleId="Normal"',
                        replacement,
                        1,
                    )
                target.writestr(item, content)
        with self.assertRaisesRegex(ValueError, "样式标识"):
            validate_template(output.getvalue(), "unsafe-style.docx")

    def test_extracts_xlsx_with_sheet_and_cell_evidence(self):
        content = (
            SAMPLES / "UC6B_5_Scope, Service Components & Pricing Breakdown.xlsx"
        ).read_bytes()
        material = validate_material(content, "pricing.xlsx")
        markdown, units = extract_xlsx_markdown(material.content, material.filename)
        self.assertEqual(1, units)
        self.assertIn("Service Component,Included,KPIs", markdown)
        self.assertIn("A2: Remote Helpdesk (8×5)", markdown)
        self.assertIn("F5: 10000", markdown)

    def test_extracts_docx_and_csv_materials_with_source_locations(self):
        docx = io.BytesIO()
        document = Document()
        document.add_paragraph("Provider CR Number: 2897654")
        document.save(docx)
        docx_material = validate_material(docx.getvalue(), "intake.docx")
        docx_markdown, docx_units = extract_docx_markdown(
            docx_material.content, docx_material.filename
        )
        csv_content = b"service,price\nHelpdesk,HKD 20000 per month\n"
        csv_material = validate_material(csv_content, "pricing.csv")
        csv_markdown, csv_units = extract_csv_markdown(
            csv_material.content, csv_material.filename
        )
        self.assertEqual("docx", docx_material.file_type)
        self.assertEqual(1, docx_units)
        self.assertIn("paragraph:1", docx_markdown)
        self.assertEqual("csv", csv_material.file_type)
        self.assertEqual(2, csv_units)
        self.assertIn("row:2", csv_markdown)

    def test_image_is_wrapped_as_single_page_pdf(self):
        content = (SAMPLES / "UC6B_9_Teams_chatlog.jfif").read_bytes()
        material = validate_material(content, "chat.jfif")
        pdf = image_to_pdf(material.content)
        document = fitz.open(stream=pdf, filetype="pdf")
        try:
            self.assertEqual(1, document.page_count)
            self.assertGreater(document[0].rect.width, 0)
        finally:
            document.close()

    def test_wrapped_images_are_submitted_to_ocr_with_pdf_filename(self):
        self.assertEqual("chat.pdf", ocr_submission_filename("chat.jfif", "image"))
        self.assertEqual(
            "evidence.pdf",
            ocr_submission_filename("evidence.pdf", "pdf"),
        )

    def test_agent_l_payloads_have_exact_two_phase_contract(self):
        template = validate_template(TEMPLATE.read_bytes(), TEMPLATE.name)
        manifest = extract_template_manifest(template.content)
        first = build_agent_l_template_payload("conv", template, manifest)
        second = build_agent_l_fill_payload(
            "conv",
            {"template_language": "en", "fields": manifest["fields"]},
            {"sources": []},
        )
        first_content = first["messages"][0]["content"]
        second_content = second["messages"][0]["content"]
        self.assertEqual(["text", "document"], [item["type"] for item in first_content])
        self.assertEqual(["text", "document"], [item["type"] for item in second_content])
        self.assertIn("original agreement template", first_content[0]["text"])
        self.assertIn("case6b_template_context.md", first_content[0]["text"])
        self.assertIn("first document attachment", first_content[0]["text"])
        self.assertIn("second document attachment", first_content[0]["text"])
        self.assertIn("Do not rely on the filenames shown", first_content[0]["text"])
        self.assertIn("return exactly one JSON Object", first_content[0]["text"])
        self.assertIn("case6b_field_fill_context.md", second_content[0]["text"])
        self.assertIn("only document attachment", second_content[0]["text"])
        self.assertIn("Do not rely on the filename shown", second_content[0]["text"])
        self.assertIn("complete field list", second_content[0]["text"])
        self.assertIn("return exactly one JSON Object", second_content[0]["text"])
        self.assertEqual(
            ["docx", "md"],
            [item["format"] for item in first_content[1]["document"]],
        )
        self.assertEqual(
            ["md"],
            [item["format"] for item in second_content[1]["document"]],
        )
        template_context = base64.b64decode(
            first_content[1]["document"][1]["base64_content"]
        ).decode("utf-8")
        fill_context = base64.b64decode(
            second_content[1]["document"][0]["base64_content"]
        ).decode("utf-8")
        self.assertIn("[CASE6B_PHASE:TEMPLATE_PARSE]", template_context)
        self.assertIn("[CASE6B_PHASE:FIELD_FILL]", fill_context)
        self.assertIn("logical_filename: `case6b_template_context.md`", template_context)
        self.assertIn("document_role: `template_context`", template_context)
        self.assertIn(
            "logical_filename: `case6b_field_fill_context.md`",
            fill_context,
        )
        self.assertIn("document_role: `field_fill_context`", fill_context)
        self.assertIn("## Detected fields", template_context)
        self.assertIn("## Evidence summary", fill_context)
        self.assertTrue(first["conversation_config"]["short_term_memory"])
        self.assertTrue(second["conversation_config"]["short_term_memory"])

    def test_needs_confirmation_preserves_suggestion_separately(self):
        manifest = {"fields": [{"field_id": "p001_f01", "field_kind": "scalar"}]}
        values = _validate_agent_fields(
            manifest,
            {
                "fields": [
                    {
                        "field_id": "p001_f01",
                        "status": "NEEDS_CONFIRMATION",
                        "value": "1 August 2026",
                        "evidence": [{"source": "email.pdf", "fact": "Proposed date"}],
                    }
                ]
            },
            filled=True,
        )
        self.assertEqual("", values[0]["value"])
        self.assertEqual("1 August 2026", values[0]["suggested_value"])

    def test_cancelled_analysis_discards_partial_results_and_stores(self):
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[MaterialRecord("m1", "facts.pdf", "pdf", b"%PDF", 1)],
            field_manifest=extract_template_manifest(TEMPLATE.read_bytes()),
        )
        local_store = {"session": [session]}
        local_results = {}
        run_id = begin_analysis_run(session)

        async def slow_summary(_material):
            await asyncio.sleep(60)
            return []

        async def scenario():
            handler = Case6BDraftingHandler()

            async def consume():
                async for _ in handler.analyze("session", local_store, local_results):
                    pass

            task = asyncio.create_task(consume())
            await asyncio.sleep(0)
            await asyncio.sleep(0)
            request_analysis_cancel(session, run_id)
            with self.assertRaises(asyncio.CancelledError):
                await task

        with patch(
            "cases.case6b_service_agreement._summarize_material",
            new=slow_summary,
        ):
            asyncio.run(scenario())
        self.assertEqual("cancelled", session.analysis_state)
        self.assertEqual({}, local_store)
        self.assertEqual({}, local_results)
        self.assertEqual([], session.fields)
        self.assertEqual("pending", session.materials[0].status)

    def test_agent_l_failure_marks_analysis_failed_and_allows_retry(self):
        material = MaterialRecord("m1", "facts.pdf", "pdf", b"%PDF", 1)
        material.status = "complete"
        material.facts = [{"fact": "Provider is Example Limited"}]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[material],
            field_manifest=extract_template_manifest(TEMPLATE.read_bytes()),
        )
        local_store = {"session": [session]}
        local_results = {}
        first_run_id = begin_analysis_run(session)

        async def scenario():
            handler = Case6BDraftingHandler()
            with self.assertRaisesRegex(RuntimeError, "Agent L failed"):
                async for _ in handler.analyze("session", local_store, local_results):
                    pass

        with patch(
            "cases.case6b_service_agreement._run_agent_l",
            new=AsyncMock(side_effect=RuntimeError("Agent L failed")),
        ):
            asyncio.run(scenario())

        self.assertEqual("failed", session.analysis_state)
        self.assertIsNone(session.active_task)
        self.assertEqual({}, local_results)
        retry_run_id = begin_analysis_run(session)
        self.assertNotEqual(first_run_id, retry_run_id)
        self.assertEqual("running", session.analysis_state)

    def test_agent_json_parser_accepts_fenced_object(self):
        parsed = parse_agent_json('```json\n{"fields": [], "conflicts": []}\n```')
        self.assertEqual([], parsed["fields"])

    def test_successful_send_prefers_new_assistant_message_over_blocking_debug_text(self):
        response = MagicMock()
        response.status_code = 200
        response.json.return_value = {
            "message": "{'answer': 'attachment debug output'}"
        }
        response.raise_for_status.return_value = None
        client = MagicMock()
        client.post = AsyncMock(return_value=response)
        client_context = MagicMock()
        client_context.__aenter__ = AsyncMock(return_value=client)
        client_context.__aexit__ = AsyncMock(return_value=False)
        with (
            patch(
                "cases.case6b_service_agreement.httpx.AsyncClient",
                return_value=client_context,
            ),
            patch(
                "cases.case6b_service_agreement._conversation_detail",
                new=AsyncMock(
                    side_effect=[
                        ("old-message", '{"old":true}'),
                        ("new-message", '{"fields":[]}'),
                    ]
                ),
            ),
        ):
            reply = asyncio.run(
                _send_message(
                    {"Authorization": "Bearer test"},
                    "conversation",
                    {"messages": []},
                )
            )
        self.assertEqual('{"fields":[]}', reply)

    def test_unknown_message_baseline_never_selects_an_old_assistant_reply(self):
        response = MagicMock()
        response.status_code = 200
        response.json.return_value = {"answer": '{"fields":[]}'}
        response.raise_for_status.return_value = None
        client = MagicMock()
        client.post = AsyncMock(return_value=response)
        client_context = MagicMock()
        client_context.__aenter__ = AsyncMock(return_value=client)
        client_context.__aexit__ = AsyncMock(return_value=False)
        with (
            patch(
                "cases.case6b_service_agreement.httpx.AsyncClient",
                return_value=client_context,
            ),
            patch(
                "cases.case6b_service_agreement._conversation_detail",
                new=AsyncMock(
                    side_effect=[
                        (None, ""),
                        ("old-message", '{"old":true}'),
                    ]
                ),
            ),
        ):
            reply = asyncio.run(
                _send_message(
                    {"Authorization": "Bearer test"},
                    "conversation",
                    {"messages": []},
                )
            )
        self.assertEqual('{"fields":[]}', reply)

    def test_review_edit_invalidates_generated_files_and_increments_version(self):
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[
                MaterialRecord("m1", "facts.pdf", "pdf", b"%PDF-1.4", 1)
            ],
            field_manifest=json.loads(
                (FIXTURES / "template_field_manifest.json").read_text("utf-8")
            ),
            fields=fields,
            review_version=1,
            generated_docx=b"old",
            generated_pdf=b"old",
        )
        updated = apply_review_changes(
            session,
            version=1,
            field_updates=[{"field_id": "p002_f01", "value": "1 August 2026"}],
            service_rows=[],
        )
        self.assertEqual(2, updated)
        changed = next(f for f in session.fields if f["field_id"] == "p002_f01")
        self.assertEqual("USER_CONFIRMED", changed["status"])
        self.assertEqual("1 August 2026", changed["value"])
        self.assertIsNone(session.generated_docx)
        self.assertIsNone(session.generated_pdf)

    def test_unchanged_review_value_keeps_agent_evidence(self):
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[],
            field_manifest=json.loads(
                (FIXTURES / "template_field_manifest.json").read_text("utf-8")
            ),
            fields=fields,
            review_version=1,
        )
        target = next(field for field in fields if field["field_id"] == "p002_f02")
        evidence = list(target["evidence"])
        apply_review_changes(
            session,
            version=1,
            field_updates=[
                {"field_id": "p002_f02", "value": target["value"]}
            ],
            service_rows=[],
        )
        self.assertEqual("FILLED", target["status"])
        self.assertEqual(evidence, target["evidence"])

    def test_review_can_accept_unchanged_agent_suggestion(self):
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[],
            field_manifest=json.loads(
                (FIXTURES / "template_field_manifest.json").read_text("utf-8")
            ),
            fields=fields,
            review_version=1,
        )
        target = next(field for field in fields if field["field_id"] == "p002_f02")
        apply_review_changes(
            session,
            version=1,
            field_updates=[],
            service_rows=[],
            accepted_field_ids=[target["field_id"]],
        )
        self.assertEqual("USER_CONFIRMED", target["status"])
        self.assertEqual("user_confirmed", target["source_type"])

    def test_review_change_invalidates_generated_files_but_keeps_document_shell(self):
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[],
            field_manifest=json.loads(
                (FIXTURES / "template_field_manifest.json").read_text("utf-8")
            ),
            fields=fields,
            review_version=1,
            document_shell=b"shell-docx",
            document_shell_hash="a" * 64,
            generated_docx=b"draft-docx",
            generated_pdf=b"%PDF-draft",
        )
        apply_review_changes(
            session,
            version=1,
            field_updates=[{"field_id": "p002_f01", "value": "1 August 2026"}],
            service_rows=[],
        )
        self.assertEqual(b"shell-docx", session.document_shell)
        self.assertEqual("a" * 64, session.document_shell_hash)
        self.assertIsNone(session.generated_docx)
        self.assertIsNone(session.generated_pdf)
    def test_docx_fill_preserves_signature_blanks_and_removes_unused_service_rows(self):
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        output = render_draft_docx(
            TEMPLATE.read_bytes(),
            json.loads(
                (FIXTURES / "template_field_manifest.json").read_text("utf-8")
            ),
            fields,
            "en",
        )
        document = Document(io.BytesIO(output))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("ServiceStar Solutions Limited", text)
        self.assertIn("Optional On-site Support", text)
        self.assertEqual(4, text.count("[TO BE CONFIRMED]"))
        self.assertEqual(6, text.count("(Price:"))
        signature_text = "\n".join(
            cell.text for table in document.tables for row in table.rows for cell in row.cells
        )
        self.assertIn("Signature:_________________________", signature_text)

    def test_conflict_resolution_requires_unique_editable_target(self):
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[],
            field_manifest=json.loads(
                (FIXTURES / "template_field_manifest.json").read_text("utf-8")
            ),
            fields=fields,
            conflicts=[
                {
                    "conflict_id": "unknown-1",
                    "semantic_key": "not_in_manifest",
                    "candidates": [{"value": "A"}, {"value": "B"}],
                    "resolved": False,
                }
            ],
            review_version=1,
        )
        with self.assertRaisesRegex(ValueError, "唯一对应"):
            apply_review_changes(
                session,
                version=1,
                field_updates=[],
                service_rows=[],
                conflict_resolutions=[
                    {"conflict_id": "unknown-1", "value": "A"}
                ],
            )
        self.assertFalse(session.conflicts[0]["resolved"])

    def test_review_supports_optional_and_blank_service_rows(self):
        manifest = json.loads(
            (FIXTURES / "template_field_manifest.json").read_text("utf-8")
        )
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[],
            field_manifest=manifest,
            fields=fields,
            review_version=1,
        )
        apply_review_changes(
            session,
            version=1,
            field_updates=[],
            service_rows=[
                {
                    "group_index": 5,
                    "action": "optional",
                    "name": "Optional On-site Support",
                    "price": "HKD 950 per hour (minimum 2 hours)",
                },
                {
                    "group_index": 7,
                    "action": "blank",
                    "name": "",
                    "price": "",
                },
            ],
        )
        row5 = [
            item
            for item in session.fields
            if next(
                field
                for field in manifest["fields"]
                if field["field_id"] == item["field_id"]
            ).get("group_index")
            == 5
        ]
        row7 = [
            item
            for item in session.fields
            if next(
                field
                for field in manifest["fields"]
                if field["field_id"] == item["field_id"]
            ).get("group_index")
            == 7
        ]
        self.assertTrue(all(item["service_action"] == "optional" for item in row5))
        self.assertTrue(all(item["status"] == "KEEP_BLANK" for item in row7))
        self.assertTrue(all(item["source_type"] == "user_confirmed" for item in row5))

    def test_conflict_resolution_updates_matching_review_field(self):
        manifest = json.loads(
            (FIXTURES / "template_field_manifest.json").read_text("utf-8")
        )
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[],
            field_manifest=manifest,
            fields=fields,
            conflicts=[
                {
                    "conflict_id": "invoice-days",
                    "semantic_key": "invoice_payment_days",
                    "candidates": [{"value": "30"}, {"value": "45"}],
                    "resolved": False,
                }
            ],
            review_version=1,
        )
        apply_review_changes(
            session,
            version=1,
            field_updates=[],
            service_rows=[],
            conflict_resolutions=[
                {"conflict_id": "invoice-days", "value": "45"}
            ],
        )
        target = next(
            field for field in session.fields if field["field_id"] == "p021_f01"
        )
        self.assertEqual("USER_CONFIRMED", target["status"])
        self.assertEqual("45", target["value"])
        self.assertTrue(session.conflicts[0]["resolved"])

    def test_conflicts_only_compare_the_same_explicit_scalar_semantics(self):
        sources = [
            {
                "facts": [
                    {
                        "semantic_key": "provider_identity",
                        "normalized_value": "ServiceStar Solutions Limited",
                        "raw_value": "ServiceStar Solutions Limited",
                        "category": "party",
                    },
                    {
                        "semantic_key": "service_price",
                        "normalized_value": "HKD 20,000 per month",
                        "raw_value": "HKD 20,000 per month",
                        "category": "service",
                    },
                    {
                        "semantic_key": "agreement_term",
                        "normalized_value": "12-month term",
                        "raw_value": "12-month term",
                        "category": "term",
                    },
                ]
            },
            {
                "facts": [
                    {
                        "semantic_key": "provider_identity",
                        "normalized_value": "ServiceStar Solutions Limited (CR No.: 2897654)",
                        "raw_value": "ServiceStar Solutions Limited, CR 2897654",
                        "category": "party",
                    },
                    {
                        "semantic_key": "provider_identity",
                        "normalized_value": "Karen",
                        "raw_value": "Karen (ServiceStar)",
                        "category": "party",
                    },
                    {
                        "semantic_key": "service_price",
                        "normalized_value": "HKD 900 per hour",
                        "raw_value": "HKD 900 per hour",
                        "category": "service",
                    },
                    {
                        "semantic_key": "agreement_term",
                        "normalized_value": "24 months",
                        "raw_value": "24 months",
                        "category": "term",
                    },
                ]
            },
        ]
        conflicts = detect_fact_conflicts(sources)
        self.assertEqual(1, len(conflicts))
        self.assertEqual("agreement_term", conflicts[0]["semantic_key"])

    def test_agent_missing_field_notes_are_not_treated_as_conflicts(self):
        normalized = normalize_agent_conflicts(
            [
                {
                    "field_id": "p002_f01",
                    "reason": "Effective date is not supplied.",
                },
                {
                    "semantic_key": "agreement_term",
                    "candidates": [{"value": "12 months"}, {"value": "24 months"}],
                },
            ]
        )
        self.assertEqual(1, len(normalized))
        self.assertEqual("agreement_term", normalized[0]["semantic_key"])
        self.assertTrue(normalized[0]["conflict_id"].startswith("agent-"))

    def test_duration_conflicts_use_raw_unit_when_normalized_value_is_numeric(self):
        conflicts = detect_fact_conflicts(
            [
                {
                    "facts": [
                        {
                            "semantic_key": "agreement_term",
                            "raw_value": "12 months",
                            "normalized_value": "12",
                            "category": "term",
                        }
                    ]
                },
                {
                    "facts": [
                        {
                            "semantic_key": "agreement_term",
                            "raw_value": "12-month term",
                            "normalized_value": "12 months",
                            "category": "term",
                        }
                    ]
                },
            ]
        )
        self.assertEqual([], conflicts)

    def test_runtime_does_not_apply_customer_profile_enrichment(self):
        manifest = extract_template_manifest(TEMPLATE.read_bytes())
        fields = [
            {
                "field_id": definition["field_id"],
                "status": "NEEDS_CONFIRMATION",
                "value": "",
                "evidence": [],
                "source_type": "evidence",
            }
            for definition in manifest["fields"]
        ]
        facts = [
            ("provider_identity", "ServiceStar Solutions Limited"),
            ("provider_cr_number", "2897654"),
            ("client_identity", "ClientCo Limited"),
            ("client_cr_number", "3234567"),
            ("provider_address", "Unit 1201, Example Tower, Hong Kong"),
            ("client_address", "88 Client Road, Hong Kong"),
            ("invoice_payment_days", "Net 30"),
            ("agreement_term", "12"),
            ("kickoff_meeting_payment_trigger", "HKD 0"),
            ("onboarding_complete_payment_trigger", "HKD 0"),
            ("service_remote_helpdesk_name", "Remote Helpdesk (8×5)"),
            (
                "service_remote_helpdesk_scope",
                "Remote incident logging, triage and support",
            ),
            ("service_remote_helpdesk_monthly_value", "20000"),
            ("service_endpoint_patching_name", "Endpoint Patching"),
            (
                "service_endpoint_patching_scope",
                "Monthly operating system and security patching",
            ),
            ("service_endpoint_patching_monthly_value", "16000"),
            (
                "service_m365_tenant_admin_name",
                "Microsoft 365 Tenant Administration",
            ),
            (
                "service_m365_tenant_admin_scope",
                "User and licence administration",
            ),
            ("service_m365_tenant_admin_monthly_value", "22000"),
            ("service_monthly_reporting_name", "Monthly KPI Reporting"),
            (
                "service_monthly_reporting_scope",
                "Monthly service performance report",
            ),
            ("service_monthly_reporting_monthly_value", "10000"),
            ("service_onsite_support_name", "On-site Support"),
            (
                "service_onsite_support_notes",
                "HKD 900/hr (min 2 hrs)",
            ),
            ("service_asset_inventory_name", "One-off Asset Inventory"),
            (
                "service_asset_inventory_notes",
                "Optional at HKD 8,000 per run",
            ),
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[],
            field_manifest=manifest,
            fields=fields,
            full_summary={
                "sources": [
                    {
                        "filename": "customer-materials",
                        "facts": [
                            {
                                "semantic_key": key,
                                "raw_value": (
                                    "12 months"
                                    if key == "agreement_term"
                                    else value
                                ),
                                "normalized_value": value,
                                "source": "customer-materials",
                                "locator": f"fact:{index}",
                            }
                            for index, (key, value) in enumerate(facts, 1)
                        ],
                    }
                ]
            },
        )

        _apply_profile_evidence_enrichment(session)

        self.assertTrue(all(field["value"] == "" for field in session.fields))
        self.assertTrue(
            all("service_action" not in field for field in session.fields)
        )

    def test_generic_renderer_does_not_apply_customer_specific_rewrites(self):
        manifest = extract_template_manifest(TEMPLATE.read_bytes())
        values = {
            "p002_f02": "ServiceStar Solutions Limited (CR No.: 2897654)",
            "p002_f03": "Unit 1201, Example Tower, Hong Kong",
            "p002_f04": "ClientCo Limited (CR No.: 3234567)",
            "p002_f05": "88 Client Road, Hong Kong",
            "p006_f01": "Remote Helpdesk (8×5) — Remote incident logging, triage and support",
            "p006_f02": "HKD 20,000 per month",
            "p007_f01": "Endpoint Patching — Monthly operating system and security patching",
            "p007_f02": "HKD 16,000 per month",
            "p008_f01": "Microsoft 365 Tenant Administration — User and licence administration",
            "p008_f02": "HKD 22,000 per month",
            "p009_f01": "Monthly KPI Reporting — Monthly service performance report",
            "p009_f02": "HKD 10,000 per month",
            "p010_f01": "Optional On-site Support",
            "p010_f02": "HKD 900 per hour (minimum 2 hours)",
            "p011_f01": "Optional One-off Asset Inventory",
            "p011_f02": "HKD 8,000 per run",
            "p016_f01": "HKD 0",
            "p016_f02": "HKD 0",
            "p021_f01": "30",
            "p024_f01": "12 months",
            "p030_f01": "30",
            "p031_f01": "10",
        }
        fields = []
        for definition in manifest["fields"]:
            field_id = definition["field_id"]
            group = definition.get("group_index")
            if field_id in values:
                action = "optional" if group in {5, 6} else "included" if group else None
                fields.append(
                    {
                        "field_id": field_id,
                        "status": "FILLED",
                        "value": values[field_id],
                        "evidence": [{"source": "gold", "fact": values[field_id]}],
                        "service_action": action,
                        "source_type": "evidence",
                    }
                )
            elif group in {7, 8, 9, 10}:
                fields.append(
                    {
                        "field_id": field_id,
                        "status": "KEEP_BLANK",
                        "value": "",
                        "evidence": [],
                        "service_action": "blank",
                        "source_type": "template_default",
                    }
                )
            else:
                fields.append(
                    {
                        "field_id": field_id,
                        "status": "LEAVE_BLANK",
                        "value": "",
                        "evidence": [],
                        "source_type": "template_default",
                    }
                )
        output = render_draft_docx(
            TEMPLATE.read_bytes(), manifest, fields, "en"
        )
        document = Document(io.BytesIO(output))
        body = "\n".join(paragraph.text for paragraph in document.paragraphs)
        signatures = "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        self.assertIn("ServiceStar Solutions Limited (CR No.: 2897654)", body)
        self.assertIn("ClientCo Limited (CR No.: 3234567)", body)
        self.assertIn("Optional On-site Support", body)
        self.assertNotIn("HKD 68,000 per service month", body)
        self.assertIn("12 months", body)
        self.assertIn("will end on 12 months", body)
        self.assertEqual(10, body.count("(Price:"))
        self.assertNotIn("ServiceStar Solutions Limited", signatures)
        self.assertNotIn("ClientCo Limited", signatures)
        self.assertIn("Name:_________________________", signatures)


class Case6BRouteTests(unittest.TestCase):
    def setUp(self):
        session_store.clear()
        session_results_store.clear()
        session_metadata.clear()
        self.client = TestClient(app)

    def test_upload_rejects_pdf_template(self):
        response = self.client.post(
            "/case6b/session/upload",
            files=[
                ("template_file", ("template.pdf", b"%PDF-1.4", "application/pdf")),
                (
                    "material_files",
                    ("facts.pdf", b"%PDF-1.4\n%%EOF", "application/pdf"),
                ),
            ],
        )
        self.assertEqual(400, response.status_code)

    def test_upload_sample_returns_session_and_37_fields(self):
        response = self.client.post(
            "/case6b/session/upload",
            files=[
                (
                    "template_file",
                    (TEMPLATE.name, TEMPLATE.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
                ),
                (
                    "material_files",
                    (
                        "pricing.xlsx",
                        (
                            SAMPLES
                            / "UC6B_5_Scope, Service Components & Pricing Breakdown.xlsx"
                        ).read_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
            ],
        )
        self.assertEqual(200, response.status_code, response.text)
        payload = response.json()
        self.assertEqual(37, payload["placeholder_count"])
        self.assertEqual(1, len(payload["materials"]))
        self.assertTrue(payload["template"]["analysis_ready"])
        self.assertNotIn("profile_id", payload["template"])
        self.assertNotIn("defaults", payload["template"])
        self.assertEqual("case6b", session_metadata[payload["session_id"]]["case_type"])

    def test_unknown_template_preflight_is_read_only_and_automatic(self):
        template_stream = io.BytesIO()
        document = Document()
        document.add_paragraph("Agreement between [Provider_Name] and [Client_Name].")
        document.save(template_stream)
        response = self.client.post(
            "/case6b/session/upload",
            files=[
                (
                    "template_file",
                    (
                        "unknown.docx",
                        template_stream.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                ),
                (
                    "material_files",
                    (
                        "pricing.xlsx",
                        (
                            SAMPLES
                            / "UC6B_5_Scope, Service Components & Pricing Breakdown.xlsx"
                        ).read_bytes(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ),
                ),
            ],
        )
        self.assertEqual(200, response.status_code, response.text)
        session_id = response.json()["session_id"]
        self.assertTrue(response.json()["template"]["analysis_ready"])
        detail = self.client.get(f"/case6b/session/{session_id}/template")
        self.assertEqual(200, detail.status_code)
        mutation = self.client.patch(
            f"/case6b/session/{session_id}/template",
            json={"version": 1, "confirmed": True},
        )
        self.assertEqual(405, mutation.status_code)

    def test_case6b_delete_is_idempotent_and_clears_shared_stores(self):
        session_id = self._ready_session()

        response = self.client.delete(f"/case6b/session/{session_id}")
        repeated = self.client.delete(f"/case6b/session/{session_id}")

        self.assertEqual(204, response.status_code)
        self.assertEqual(204, repeated.status_code)
        self.assertNotIn(session_id, session_store)
        self.assertNotIn(session_id, session_results_store)
        self.assertNotIn(session_id, session_metadata)

    def test_case6b_delete_rejects_busy_session(self):
        session_id = self._ready_session()
        session = session_store[session_id][0]
        asyncio.run(session.lock.acquire())
        try:
            response = self.client.delete(f"/case6b/session/{session_id}")
        finally:
            session.lock.release()

        self.assertEqual(409, response.status_code)
        self.assertEqual("SESSION_BUSY", response.json()["detail"]["code"])
        self.assertIn(session_id, session_store)
        self.assertIn(session_id, session_results_store)

    def test_review_unknown_session_returns_410(self):
        response = self.client.get("/case6b/session/missing/review")
        self.assertEqual(410, response.status_code)
        self.assertEqual("SESSION_EXPIRED", response.json()["detail"]["code"])

    def _ready_session(self) -> str:
        session_id = "case6b-ready"
        manifest = json.loads(
            (FIXTURES / "template_field_manifest.json").read_text("utf-8")
        )
        fields = json.loads((FIXTURES / "expected_fill.json").read_text("utf-8"))[
            "fields"
        ]
        session = Case6BSession(
            template=validate_template(TEMPLATE.read_bytes(), TEMPLATE.name),
            materials=[],
            field_manifest=manifest,
            fields=fields,
            review_version=1,
        )
        session_store[session_id] = [session]
        session_metadata[session_id] = {"case_type": "case6b"}
        session_results_store[session_id] = {
            "case_type": "case6b",
            "results": {"协议草案": {"conversation_id": "private"}},
        }
        return session_id

    def test_review_rejects_stale_version_and_signature_edit(self):
        session_id = self._ready_session()
        stale = self.client.patch(
            f"/case6b/session/{session_id}/review",
            json={"version": 9, "field_updates": [], "service_rows": []},
        )
        self.assertEqual(409, stale.status_code)
        signature = self.client.patch(
            f"/case6b/session/{session_id}/review",
            json={
                "version": 1,
                "field_updates": [
                    {
                        "field_id": "t000_r000_c000_p002_f01",
                        "value": "signed",
                    }
                ],
                "service_rows": [],
            },
        )
        self.assertEqual(400, signature.status_code)

    def test_finalize_requires_unresolved_confirmation_and_keeps_docx_on_pdf_failure(self):
        session_id = self._ready_session()
        unresolved = next(
            field
            for field in session_store[session_id][0].fields
            if field["field_id"] == "p016_f01"
        )
        unresolved.update({"status": "NEEDS_CONFIRMATION", "value": "", "evidence": []})
        blocked = self.client.post(
            f"/case6b/session/{session_id}/finalize",
            json={"version": 1, "allow_unresolved": False},
        )
        self.assertEqual(409, blocked.status_code)
        with patch(
            "api.case6b_routes.convert_docx_to_pdf",
            side_effect=RuntimeError("PDF unavailable"),
        ):
            finalized = self.client.post(
                f"/case6b/session/{session_id}/finalize",
                json={"version": 1, "allow_unresolved": True},
            )
        self.assertEqual(200, finalized.status_code, finalized.text)
        self.assertTrue(finalized.json()["docx_ready"])
        self.assertFalse(finalized.json()["pdf_ready"])
        report = self.client.post(
            "/pdf/session/report",
            data={
                "session_id": session_id,
                "party": "协议草案",
                "output_format": "docx",
            },
        )
        self.assertEqual(200, report.status_code)
        self.assertTrue(report.content.startswith(b"PK"))

    def test_document_view_returns_word_shell_with_unique_field_markers(self):
        session_id = self._ready_session()
        view = self.client.get(f"/case6b/session/{session_id}/document-view")
        self.assertEqual(200, view.status_code, view.text)
        payload = view.json()
        self.assertEqual(37, len(payload["fields"]))
        shell = self.client.get(payload["shell_url"])
        self.assertEqual(200, shell.status_code, shell.text)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            shell.headers["content-type"],
        )
        document = Document(io.BytesIO(shell.content))
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [
                paragraph.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
                for paragraph in cell.paragraphs
            ]
        )
        for field in payload["fields"]:
            self.assertEqual(1, text.count(field["marker"]), field["field_id"])

    def test_document_shell_rejects_stale_manifest_hash(self):
        session_id = self._ready_session()
        response = self.client.get(
            f"/case6b/session/{session_id}/document-shell"
            f"?manifest_hash={'0' * 64}"
        )
        self.assertEqual(409, response.status_code)
        self.assertEqual("STALE_DOCUMENT_VIEW", response.json()["detail"]["code"])

    def test_cancel_is_idempotent_and_rejects_stale_run(self):
        session_id = self._ready_session()
        session = session_store[session_id][0]
        session.analysis_run_id = "run-current-1234"
        session.analysis_state = "running"
        stale = self.client.post(
            f"/case6b/session/{session_id}/cancel",
            json={"run_id": "run-stale-12345"},
        )
        self.assertEqual(409, stale.status_code)
        cancelled = self.client.post(
            f"/case6b/session/{session_id}/cancel",
            json={"run_id": "run-current-1234"},
        )
        repeated = self.client.post(
            f"/case6b/session/{session_id}/cancel",
            json={"run_id": "run-current-1234"},
        )
        self.assertEqual(202, cancelled.status_code)
        self.assertEqual(202, repeated.status_code)
        self.assertTrue(session.cancel_event.is_set())
        self.assertEqual("cancelling", session.analysis_state)

    def test_finalize_renders_current_review_without_pdf_preview_cache(self):
        session_id = self._ready_session()
        session = session_store[session_id][0]
        with (
            patch("api.case6b_routes.render_draft_docx", return_value=b"draft-docx") as render,
            patch("api.case6b_routes.convert_docx_to_pdf", return_value=b"%PDF-draft") as convert,
        ):
            finalized = self.client.post(
                f"/case6b/session/{session_id}/finalize",
                json={"version": 1, "allow_unresolved": True},
            )
        self.assertEqual(200, finalized.status_code, finalized.text)
        self.assertEqual(b"draft-docx", session.generated_docx)
        self.assertEqual(b"%PDF-draft", session.generated_pdf)
        render.assert_called_once()
        convert.assert_called_once()
    def test_finalize_is_blocked_by_unresolved_conflicts(self):
        session_id = self._ready_session()
        session_store[session_id][0].conflicts = [
            {
                "conflict_id": "term-1",
                "semantic_key": "agreement_term",
                "candidates": [{"value": "12 months"}, {"value": "24 months"}],
                "resolved": False,
            }
        ]
        blocked = self.client.post(
            f"/case6b/session/{session_id}/finalize",
            json={"version": 1, "allow_unresolved": True},
        )
        self.assertEqual(409, blocked.status_code)
        self.assertEqual("UNRESOLVED_CONFLICTS", blocked.json()["detail"]["code"])

    def test_single_material_retry_rejects_unknown_or_nonfailed_target(self):
        session_id = self._ready_session()
        session = session_store[session_id][0]
        session.materials = [
            validate_material(
                (SAMPLES / "UC6B_6_Heads of Terms.pdf").read_bytes(),
                "UC6B_6_Heads of Terms.pdf",
            )
        ]
        unknown = self.client.post(
            f"/case6b/session/{session_id}/retry?material_id=missing"
        )
        self.assertEqual(404, unknown.status_code)
        not_failed = self.client.post(
            f"/case6b/session/{session_id}/retry"
            f"?material_id={session.materials[0].material_id}"
        )
        self.assertEqual(400, not_failed.status_code)
        self.assertEqual(
            "MATERIAL_NOT_FAILED",
            not_failed.json()["detail"]["code"],
        )


if __name__ == "__main__":
    unittest.main()
