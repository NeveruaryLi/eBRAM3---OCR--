"""Run the real Case 6B customer sample through the local application."""

from __future__ import annotations

import argparse
import io
import json
import sys
from pathlib import Path

from docx import Document
from dotenv import load_dotenv


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--env-file", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    load_dotenv(args.env_file, override=True)
    root = Path(__file__).resolve().parents[1]
    sys.path.insert(0, str(root))

    from fastapi.testclient import TestClient

    from app import app

    source = root / "input_example" / "Usecase 6B" / "Input documents"
    template = source / "UC6B_0_Format - Service Agreement.docx"
    materials = sorted(path for path in source.iterdir() if path != template)
    files = [
        (
            "template_file",
            (
                template.name,
                template.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )
    ]
    files.extend(
        ("material_files", (path.name, path.read_bytes(), "application/octet-stream"))
        for path in materials
    )
    with TestClient(app) as client:
        upload = client.post("/case6b/session/upload", files=files)
        upload.raise_for_status()
        uploaded = upload.json()
        print(
            json.dumps(
                {
                    "phase": "upload",
                    "profile_id": uploaded["template"]["profile_id"],
                    "confirmed": uploaded["template"]["confirmed"],
                    "materials": len(uploaded["materials"]),
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
        session_id = uploaded["session_id"]
        analysis = client.post(
            "/pdf/session/analyze", data={"session_id": session_id}
        )
        analysis.raise_for_status()
        events = [
            json.loads(line[5:].strip())
            for line in analysis.text.splitlines()
            if line.startswith("data:")
        ]
        errors = [
            event
            for event in events
            if event.get("type") in {"error", "fatal_error"}
            and event.get("material_id")
        ]
        for _ in range(2):
            if not errors:
                break
            retry_events = []
            for failure in errors:
                retry = client.post(
                    f"/case6b/session/{session_id}/retry",
                    params={"material_id": failure["material_id"]},
                )
                retry.raise_for_status()
                retry_events.extend(
                    json.loads(line[5:].strip())
                    for line in retry.text.splitlines()
                    if line.startswith("data:")
                )
            events.extend(retry_events)
            errors = [
                event
                for event in retry_events
                if event.get("type") in {"error", "fatal_error"}
                and event.get("material_id")
            ]
        print(
            json.dumps(
                {
                    "phase": "analysis",
                    "events": len(events),
                    "errors": errors,
                    "result": next(
                        (event for event in events if event.get("type") == "result"),
                        None,
                    ),
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
        if errors:
            return 2
        review_response = client.get(f"/case6b/session/{session_id}/review")
        review_response.raise_for_status()
        review = review_response.json()
        fields = review["fields"]
        actions: dict[str, int] = {}
        for field in fields:
            if field.get("field_kind") != "repeatable_service":
                continue
            action = field.get("service_action", "included")
            actions[action] = actions.get(action, 0) + 1
        key_values = {
            field.get("semantic_key"): field.get("value")
            for field in fields
            if field.get("semantic_key")
            in {
                "provider_identity",
                "client_identity",
                "agreement_term",
                "termination_notice",
                "materials_return",
            }
        }
        print(
            json.dumps(
                {
                    "phase": "review",
                    "fields": len(fields),
                    "service_actions": actions,
                    "unresolved_fields": review["unresolved_count"],
                    "unresolved_conflicts": review["unresolved_conflict_count"],
                    "key_values": key_values,
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
        if review["unresolved_conflict_count"]:
            return 3
        finalized = client.post(
            f"/case6b/session/{session_id}/finalize",
            json={"version": review["version"], "allow_unresolved": True},
        )
        finalized.raise_for_status()
        result = finalized.json()
        report = client.post(
            "/pdf/session/report",
            data={
                "session_id": session_id,
                "party": "协议草案",
                "output_format": "docx",
            },
        )
        report.raise_for_status()
        args.output_dir.mkdir(parents=True, exist_ok=True)
        output_path = args.output_dir / "case6b-v1.1-customer-sample.docx"
        output_path.write_bytes(report.content)
        pdf_ready = False
        if result["pdf_ready"]:
            pdf_report = client.post(
                "/pdf/session/report",
                data={
                    "session_id": session_id,
                    "party": "协议草案",
                    "output_format": "pdf",
                },
            )
            pdf_report.raise_for_status()
            pdf_ready = pdf_report.content.startswith(b"%PDF-")
            (args.output_dir / "case6b-v1.1-customer-sample.pdf").write_bytes(
                pdf_report.content
            )
        document = Document(io.BytesIO(report.content))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        text += "\n" + "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        checks = {
            "provider_cr": "2897654" in text,
            "client_cr": "3234567" in text,
            "monthly_total": "68,000" in text,
            "zero_signing_and_onboarding": text.count("HKD 0") >= 2,
            "optional_services": all(
                value in text
                for value in (
                    "On-site Support",
                    "Asset Inventory",
                )
            ),
            "optional_prices": all(
                value in text
                for value in ("900 per hour", "8,000 per run")
            ),
            "ten_service_rows": text.count("(Price:") == 10,
            "term_12_months": "12 months" in text or "one year" in text.lower(),
            "termination_30": "30 days" in text,
            "return_10": "10 days" in text,
            "signature_companies": all(
                value in text
                for value in (
                    "ServiceStar Solutions Limited",
                    "ClientCo Limited",
                )
            ),
            "pdf_download": pdf_ready,
        }
        print(
            json.dumps(
                {
                    "phase": "finalize",
                    "docx_ready": result["docx_ready"],
                    "pdf_ready": result["pdf_ready"],
                    "checks": checks,
                    "output": str(output_path),
                },
                ensure_ascii=False,
            ),
            flush=True,
        )
        return 0 if all(checks.values()) else 4


if __name__ == "__main__":
    raise SystemExit(main())
