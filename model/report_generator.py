"""
报告生成工具 v2

将 Agent B 的 Markdown 格式综合分析文本转换为格式化的 Word/PDF 文档。
报告内容仅包含 Agent B 的回复，不包含原始材料数据。
"""

from __future__ import annotations

import logging
import os
import re
import tempfile

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

logger = logging.getLogger(__name__)

# ── 字体常量 ───────────────────────────────────────────────────────────────────
_FONT_CN   = "微软雅黑"
_FONT_CODE = "Courier New"

# Heading sizes (pt)
_H_SIZES = {1: 16, 2: 14, 3: 12}
_BODY_SIZE = 11.0

# Line spacing multiplier (1.5x) expressed in twips: 1.5 × 240 = 360
_LINE_SPACING_TWIPS = "360"
# Paragraph spacing in twips: 6pt × 20 = 120
_SPACE_6PT  = "120"
_SPACE_4PT  = "80"
_SPACE_10PT = "200"
_SPACE_2PT  = "40"


# ── XML / font helpers ─────────────────────────────────────────────────────────

def _set_run_font(
    run,
    name: str = _FONT_CN,
    size_pt: float | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
    color_rgb: tuple | None = None,
) -> None:
    """Set font name (including East Asian), size, weight and colour on a run."""
    run.font.name = name
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color_rgb is not None:
        run.font.color.rgb = RGBColor(*color_rgb)

    # East Asian font must be set via raw XML
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"),   name)
    rFonts.set(qn("w:hAnsi"),   name)
    rFonts.set(qn("w:eastAsia"), name)


def _set_para_spacing(
    para,
    before_twips: str = _SPACE_6PT,
    after_twips: str  = _SPACE_6PT,
    line_twips: str   = _LINE_SPACING_TWIPS,
) -> None:
    """Set paragraph spacing and line height via XML (works across all styles)."""
    pPr = para._element.get_or_add_pPr()
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"),   before_twips)
    spacing.set(qn("w:after"),    after_twips)
    spacing.set(qn("w:line"),     line_twips)
    spacing.set(qn("w:lineRule"), "auto")


def _apply_style_font(doc: Document, style_name: str, font_name: str = _FONT_CN) -> None:
    """Patch the East Asian font and force black colour on a named paragraph style."""
    try:
        style = doc.styles[style_name]
    except KeyError:
        return
    style.font.name = font_name
    # Force black — overrides Word's built-in theme colour (e.g. blue on Heading styles)
    style.font.color.rgb = RGBColor(0, 0, 0)

    rPr = style.element.get_or_add_rPr()

    # Font name (East Asian)
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"),    font_name)
    rFonts.set(qn("w:hAnsi"),    font_name)
    rFonts.set(qn("w:eastAsia"), font_name)

    # Colour — remove any theme-colour node and write a solid black value
    for old in rPr.findall(qn("w:color")):
        rPr.remove(old)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), "000000")
    rPr.append(color_el)


def _configure_doc(doc: Document) -> None:
    """Apply base font to all styles we use so auto-generated runs inherit it."""
    for sname in ("Normal", "Heading 1", "Heading 2", "Heading 3",
                  "List Bullet", "List Number"):
        _apply_style_font(doc, sname)
    # Page margins: left/right ≈ 3.17 cm, top/bottom ≈ 2.54 cm (in EMU)
    sec = doc.sections[0]
    sec.top_margin    = Pt(72)
    sec.bottom_margin = Pt(72)
    sec.left_margin   = Pt(90)
    sec.right_margin  = Pt(90)


# ── Inline Markdown parser ─────────────────────────────────────────────────────

_INLINE_RE = re.compile(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`")


def _add_inline_runs(para, text: str, size_pt: float = _BODY_SIZE) -> None:
    """
    Append runs to *para* with bold/italic/code inline Markdown resolved.
    Handles: **bold**, *italic*, `inline code`.
    """
    last = 0
    for m in _INLINE_RE.finditer(text):
        # Text before this match
        if m.start() > last:
            run = para.add_run(text[last:m.start()])
            _set_run_font(run, size_pt=size_pt)
        # Bold
        if m.group(1) is not None:
            run = para.add_run(m.group(1))
            _set_run_font(run, size_pt=size_pt, bold=True)
        # Italic
        elif m.group(2) is not None:
            run = para.add_run(m.group(2))
            _set_run_font(run, size_pt=size_pt, italic=True)
        # Inline code
        elif m.group(3) is not None:
            run = para.add_run(m.group(3))
            _set_run_font(run, name=_FONT_CODE, size_pt=size_pt - 0.5)
        last = m.end()

    # Remaining text
    if last < len(text):
        run = para.add_run(text[last:])
        _set_run_font(run, size_pt=size_pt)


# ── Block-level builders ──────────────────────────────────────────────────────

def _add_heading(doc: Document, text: str, level: int) -> None:
    size_pt  = _H_SIZES.get(level, 12)
    before   = _SPACE_10PT if level == 1 else _SPACE_6PT
    after    = _SPACE_4PT

    style_name = f"Heading {min(level, 3)}"
    try:
        para = doc.add_paragraph(style=style_name)
    except KeyError:
        para = doc.add_paragraph()

    run = para.add_run(text.strip())
    _set_run_font(run, size_pt=size_pt, bold=True, color_rgb=(0, 0, 0))
    _set_para_spacing(para, before_twips=before, after_twips=after)


def _add_body(doc: Document, text: str) -> None:
    para = doc.add_paragraph()
    _add_inline_runs(para, text, size_pt=_BODY_SIZE)
    _set_para_spacing(para)


def _add_bullet(doc: Document, text: str, ordered: bool = False) -> None:
    style = "List Number" if ordered else "List Bullet"
    try:
        para = doc.add_paragraph(style=style)
    except KeyError:
        para = doc.add_paragraph()
        run = para.add_run("• " if not ordered else "  ")
        _set_run_font(run, size_pt=_BODY_SIZE)

    _add_inline_runs(para, text, size_pt=_BODY_SIZE)
    _set_para_spacing(para, before_twips=_SPACE_2PT, after_twips=_SPACE_2PT)


def _add_code_block(doc: Document, lines: list) -> None:
    text = "\n".join(lines)
    para = doc.add_paragraph()
    run  = para.add_run(text)
    _set_run_font(run, name=_FONT_CODE, size_pt=10.0)

    # Light grey shading
    pPr = para._element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  "F0F0F0")
    pPr.append(shd)
    _set_para_spacing(para, before_twips=_SPACE_4PT, after_twips=_SPACE_4PT,
                      line_twips="288")  # 1.2 × 240


_ORDERED_RE = re.compile(r"^\d+[.)]\s+")


def _markdown_to_doc(doc: Document, md_text: str) -> None:
    """Convert Markdown text to Word paragraphs line by line."""
    lines = md_text.split("\n")
    i = 0
    n = len(lines)

    while i < n:
        raw     = lines[i]
        stripped = raw.strip()

        # ── Fenced code block ─────────────────────────────────────────
        if stripped.startswith("```"):
            code_lines: list = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            _add_code_block(doc, code_lines)
            i += 1          # skip closing ```
            continue

        # ── ATX headings ──────────────────────────────────────────────
        if stripped.startswith("### "):
            _add_heading(doc, stripped[4:], 3)
        elif stripped.startswith("## "):
            _add_heading(doc, stripped[3:], 2)
        elif stripped.startswith("# "):
            _add_heading(doc, stripped[2:], 1)

        # ── Unordered list ────────────────────────────────────────────
        elif stripped.startswith(("- ", "* ", "• ")):
            _add_bullet(doc, stripped[2:], ordered=False)

        # ── Ordered list ──────────────────────────────────────────────
        elif _ORDERED_RE.match(stripped):
            content = _ORDERED_RE.sub("", stripped)
            _add_bullet(doc, content, ordered=True)

        # ── Horizontal rule ───────────────────────────────────────────
        elif re.match(r"^[-*_]{3,}$", stripped):
            # Add a thin border paragraph
            para = doc.add_paragraph()
            pPr  = para._element.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"),   "single")
            bottom.set(qn("w:sz"),    "4")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "CCCCCC")
            pBdr.append(bottom)
            pPr.append(pBdr)
            _set_para_spacing(para, before_twips=_SPACE_2PT, after_twips=_SPACE_2PT,
                              line_twips="240")

        # ── Empty line ────────────────────────────────────────────────
        elif not stripped:
            pass   # skip — spacing comes from paragraph format

        # ── Normal body paragraph ─────────────────────────────────────
        else:
            _add_body(doc, stripped)

        i += 1


# ── Public API ────────────────────────────────────────────────────────────────

def generate_report(
    analysis_text: str,
    output_format: str = "pdf",
    docs: list | None = None,       # 废弃参数，保留兼容旧调用签名
    base_name: str = "ebram_analysis",
) -> tuple[bytes, str]:
    """
    将 Agent B 综合分析文本（Markdown）转换为 Word/PDF 报告。

    Args:
        analysis_text: Agent B 的综合分析回复（Markdown 格式）
        output_format: "pdf"（默认）或 "docx"
        docs:          废弃参数，保留以兼容旧调用签名，不再写入报告
        base_name:     输出文件名（不含扩展名），默认 "ebram_analysis"

    Returns:
        (file_bytes, filename) 元组
    """
    doc = Document()
    _configure_doc(doc)

    if analysis_text and analysis_text.strip():
        _markdown_to_doc(doc, analysis_text)
    else:
        para = doc.add_paragraph("（综合分析尚未生成，请先点击「开始综合分析」后再下载报告）")
        for run in para.runs:
            _set_run_font(run, color_rgb=(0x80, 0x80, 0x80), size_pt=_BODY_SIZE)
        _set_para_spacing(para)

    safe_name = base_name.replace(" ", "_") or "ebram_analysis"

    with tempfile.TemporaryDirectory() as tmp_dir:
        docx_path = os.path.join(tmp_dir, f"{safe_name}.docx")
        doc.save(docx_path)
        logger.info("Word 文档已生成：%.1f KB", os.path.getsize(docx_path) / 1024)

        fmt = output_format.lower().strip()

        if fmt == "pdf":
            pdf_path = os.path.join(tmp_dir, f"{safe_name}.pdf")
            try:
                from docx2pdf import convert  # noqa: PLC0415
                convert(docx_path, pdf_path)
                logger.info("PDF 转换完成：%.1f KB", os.path.getsize(pdf_path) / 1024)
                with open(pdf_path, "rb") as f:
                    return f.read(), f"{safe_name}.pdf"
            except Exception as exc:
                logger.warning("PDF 转换失败，降级返回 docx：%s", exc)
                with open(docx_path, "rb") as f:
                    return f.read(), f"{safe_name}.docx"

        # docx
        with open(docx_path, "rb") as f:
            return f.read(), f"{safe_name}.docx"
